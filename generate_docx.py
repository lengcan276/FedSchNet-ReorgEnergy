"""
Convert paper/main.tex → paper/main.docx
Nature Machine Intelligence style — KA-GNN inspired formatting.

Features:
- Actual PNG figure insertion at correct inline positions
- Centered italic 10pt figure captions
- Word heading styles (Heading 1/2/3) with proper sizing
- Full LaTeX → Unicode conversion
- Tables from CSV / hardcoded data
- Reference list parsed from references.bib
- Zero LaTeX residuals
"""

import re
import csv
import os
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

PROJECT_ROOT = Path(__file__).resolve().parent
TEX_PATH = PROJECT_ROOT / 'paper' / 'main.tex'
BIB_PATH = PROJECT_ROOT / 'paper' / 'references.bib'
DOCX_PATH = PROJECT_ROOT / 'paper' / 'main.docx'
RESULTS_DIR = PROJECT_ROOT / 'results' / 'tables'
FIGURES_DIR = PROJECT_ROOT / 'results' / 'figures'

# ============================================================
#  FIGURE DEFINITIONS
# ============================================================

FIGURE_DEFS = {
    'fig1': {
        'file': 'fig1_overview.png',
        'width': 6.0,
        'caption': (
            'Figure 1. Graphical abstract of the FedPer-KAN framework. '
            'Four heterogeneous clients (A\u2013D) collaboratively train shared GNN encoders '
            'while maintaining private KAN regression heads for domain-specific '
            'reorganization energy prediction.'
        ),
    },
    'fig2': {
        'file': 'fig2_dataset_distribution.png',
        'width': 6.0,
        'caption': (
            'Figure 2. Dataset overview: molecular weight distributions and label (\u03bb) '
            'ranges for all four clients. Client D exhibits the broadest \u03bb_hole range '
            '(0.45\u20133.00 eV), reflecting the conformational diversity of TADF emitters.'
        ),
    },
    'fig3': {
        'file': 'fig3_fedper_workflow.png',
        'width': 6.0,
        'caption': (
            'Figure 3. FedPer workflow. Step 1: Local training with private encoder + KAN head. '
            'Step 2: Upload encoder parameters only. Step 3: FedAvg aggregation. '
            'Step 4: Download updated encoder; KAN heads remain unchanged.'
        ),
    },
    'fig4': {
        'file': 'fig4_tsne.png',
        'width': 6.0,
        'caption': (
            'Figure 4. t-SNE visualization of Morgan fingerprint chemical space. '
            'Client D (TADF) occupies a distinct region characterized by complex '
            'donor\u2013acceptor architectures. Client C (conjugated semiconductors) '
            'bridges Clients A/B and Client D.'
        ),
    },
    'fig5': {
        'file': 'fig5_kan_vs_mlp.png',
        'width': 3.3,
        'caption': (
            'Figure 5. Comparison of KAN vs. MLP regression heads under local training. '
            'KAN outperforms MLP by 30.7% in R\u00b2 on Client A (E4 vs. E1).'
        ),
    },
    'fig6': {
        'file': 'fig6_strategies_ssl.png',
        'width': 6.0,
        'caption': (
            'Figure 6. (a) Federated strategy comparison: Local, FedAvg, FedPer, '
            'FedPer+FedBN for GIN+KAN. (b) Self-supervised pre-training: '
            'AtomMask, EdgePred, GraphCL.'
        ),
    },
    'fig7': {
        'file': 'fig7_physics_ml_fusion.png',
        'width': 6.0,
        'caption': (
            'Figure 7. Physics\u2013ML fusion overview. (a) 3D conformer showing torsional '
            'degrees of freedom. (b) 2D molecular graph with topological connectivity only. '
            '(c) Representation bottleneck: R\u00b2 progression from local 2D GIN '
            'to federated 3D SchNet on rigid D\u2013A molecules.'
        ),
    },
    'fig8': {
        'file': 'fig8_waterfall.png',
        'width': 6.0,
        'caption': (
            'Figure 8. Waterfall chart showing cumulative R\u00b2 gains from each '
            'methodological component on Client D \u03bb_hole prediction.'
        ),
    },
    'fig9': {
        'file': 'fig9_ablation_combined.png',
        'width': 6.0,
        'caption': (
            'Figure 9. Ablation studies on Client D \u03bb_hole R\u00b2: '
            '(a) KAN grid size, (b) SSL pre-training rounds, '
            '(c) data size, (d) communication rounds.'
        ),
    },
    'fig13': {
        'file': 'fig13_gin_vs_schnet.png',
        'width': 6.0,
        'caption': (
            'Figure 13. Comparison of 2D GIN (E12) vs. 3D SchNet (E14) across all clients. '
            'SchNet provides the largest improvement on Client D (\u0394R\u00b2 = +0.212).'
        ),
    },
    'fig14': {
        'file': 'fig14_scatter.png',
        'width': 3.3,
        'caption': (
            'Figure 14. Scatter plot of predicted vs. true \u03bb_hole for the '
            'best-performing model (E14, SchNet+KAN FedAvg) on Client D.'
        ),
    },
    'fig15': {
        'file': 'fig15_case_study.png',
        'width': 6.0,
        'caption': (
            'Figure 15. Per-molecule case study on Client D. '
            '(a) Predicted vs. true \u03bb_hole for representative molecules. '
            '(b) Absolute error vs. number of rotatable bonds; GIN errors show '
            'significant correlation (r = 0.30, p = 0.028), SchNet errors do not.'
        ),
    },
    'fig16': {
        'file': 'fig16_representative_molecules.png',
        'width': 6.0,
        'caption': (
            'Figure 16. Ball-and-stick representations of representative Client D molecules, '
            'categorized by SchNet prediction quality: best (green), medium (gold), worst (red).'
        ),
    },
    'fig17': {
        'file': 'fig17_subgroup_analysis.png',
        'width': 6.0,
        'caption': (
            'Figure 17. Subgroup analysis of Client D predictions. '
            '(a) R\u00b2 comparison across subgroups for GIN and SchNet. '
            '(b) Predicted vs. true \u03bb_hole for rigid high-DA subgroup '
            '(n = 13, SchNet R\u00b2 = 0.852).'
        ),
    },
}

# Figures to insert BEFORE a heading whose number starts with the given prefix.
# Checked in order; most specific prefixes first to avoid ambiguity.
FIGURE_PLACEMENTS = [
    ('2. ',    ['fig1']),
    ('3.5 ',   ['fig3']),
    ('4.2 ',   ['fig2', 'fig4']),
    ('5.2 ',   ['fig5']),
    ('5.3 ',   ['fig6']),
    ('5.6 ',   ['fig9']),
    ('5.7.1 ', ['fig13', 'fig7']),
    ('5.7 ',   ['fig8']),
    ('6. ',    ['fig15', 'fig16', 'fig14', 'fig17']),
]


# ============================================================
#  BIB PARSING
# ============================================================

def _extract_brace_value(text, start):
    """Extract content between matched braces starting at position `start`.
    Returns (content, end_pos) where end_pos is the index after the closing '}'.
    """
    depth = 0
    i = start
    while i < len(text):
        if text[i] == '{':
            if depth == 0:
                content_start = i + 1
            depth += 1
        elif text[i] == '}':
            depth -= 1
            if depth == 0:
                return text[content_start:i], i + 1
        i += 1
    return text[start:], len(text)


def _clean_bib_value(fval):
    """Clean a bib field value: handle LaTeX accents and remove braces."""
    # Handle LaTeX accents BEFORE removing braces
    fval = fval.replace(r'{\"u}', 'ü').replace(r'{\"o}', 'ö').replace(r'{\"a}', 'ä')
    fval = fval.replace(r"{\'e}", 'é').replace(r"{\^o}", 'ô').replace(r"{\'i}", 'í')
    fval = fval.replace(r'{\v{c}}', 'č').replace(r'{\v c}', 'č')
    fval = fval.replace(r'{\ss}', 'ß')
    # Now handle bare accent commands (without outer braces)
    fval = fval.replace(r'\"u', 'ü').replace(r'\"o', 'ö').replace(r'\"a', 'ä')
    fval = fval.replace(r"\'e", 'é').replace(r"\^o", 'ô').replace(r"\'i", 'í')
    fval = fval.replace(r'\v{c}', 'č').replace(r"\v c", 'č')
    fval = fval.replace(r'\ss', 'ß').replace(r'\&', '&')
    # Handle \url{...} → extract content
    fval = re.sub(r'\\url\{([^}]*)\}', r'\1', fval)
    # Remove remaining braces (e.g., {T}heory → Theory, {KAN} → KAN)
    fval = fval.replace('{', '').replace('}', '')
    # Remove remaining backslash commands but not the content
    fval = fval.replace('\\', '')
    fval = re.sub(r'\s+', ' ', fval).strip()
    return fval


def parse_bib(bib_path):
    """Parse references.bib → dict of {key: {fields}}.
    Handles nested braces correctly.
    """
    entries = {}
    text = bib_path.read_text(encoding='utf-8')
    # Match entry starts: @type{key,
    entry_pattern = re.compile(r'@\w+\{(\w+),')
    for em in entry_pattern.finditer(text):
        key = em.group(1)
        body_start = em.end()
        # Find the matching closing brace for the entry
        depth = 1
        i = body_start
        while i < len(text) and depth > 0:
            if text[i] == '{':
                depth += 1
            elif text[i] == '}':
                depth -= 1
            i += 1
        body = text[body_start:i - 1]

        fields = {}
        # Parse fields: fieldname = {value} with proper brace matching
        field_pattern = re.compile(r'(\w+)\s*=\s*\{')
        for fm in field_pattern.finditer(body):
            fname = fm.group(1).lower()
            # Find the opening brace and extract balanced content
            brace_start = fm.end() - 1  # position of '{'
            fval, _ = _extract_brace_value(body, brace_start)
            fields[fname] = _clean_bib_value(fval)
        entries[key] = fields
    return entries


def make_cite_map(bib_entries):
    """Build citation key → '[Author et al., Year]' mapping."""
    cmap = {}
    for key, fields in bib_entries.items():
        author = fields.get('author', '')
        year = fields.get('year', '????')
        if ' and ' in author:
            authors = [a.strip() for a in author.split(' and ')]
            first_surname = authors[0].split(',')[0].strip()
            if len(authors) == 2:
                second_surname = authors[1].split(',')[0].strip()
                cite_str = f'{first_surname} and {second_surname}, {year}'
            else:
                cite_str = f'{first_surname} et al., {year}'
        elif author:
            first_surname = author.split(',')[0].strip()
            cite_str = f'{first_surname}, {year}'
        else:
            cite_str = f'{key}, {year}'
        cmap[key] = cite_str
    return cmap


# ============================================================
#  REF LABEL → DISPLAY MAPPING
# ============================================================

REF_MAP = {
    'fig:overview': 'Fig. 1', 'fig:workflow': 'Fig. 3',
    'fig:dataset': 'Fig. 2', 'fig:tsne': 'Fig. 4',
    'fig:kan_vs_mlp': 'Fig. 5', 'fig:strategies': 'Fig. 6',
    'fig:ablation': 'Fig. 9', 'fig:gin_vs_schnet': 'Fig. 13',
    'fig:case_study': 'Fig. 15', 'fig:representative_mols': 'Fig. 16',
    'fig:scatter': 'Fig. 14', 'fig:subgroup': 'Fig. 17', 'fig:physics_ml': 'Fig. 7',
    'tab:summary': 'Table 1', 'tab:subgroup': 'Table 2',
    'eq:fed_obj': 'Eq. 1',
    'sec:intro': 'Section 1', 'sec:related': 'Section 2',
    'sec:methods': 'Section 3', 'sec:setup': 'Section 4',
    'sec:results': 'Section 5', 'sec:conclusions': 'Section 6',
    'sec:kan_vs_mlp': 'Section 5.1', 'sec:fed_strategy': 'Section 5.2',
    'sec:ssl': 'Section 5.3', 'sec:four_client': 'Section 5.4',
    'sec:ablation': 'Section 5.5', 'sec:bottleneck': 'Section 5.6',
    'sec:schnet': 'Section 5.7', 'sec:subgroup': 'Section 5.7.1',
    'sec:formulation': 'Section 3.1', 'sec:graph': 'Section 3.2',
    'sec:architecture': 'Section 3.3', 'sec:fedstrategy': 'Section 3.4',
    'sec:labelnorm': 'Section 3.5', 'sec:datasets': 'Section 4.1',
    'sec:implementation': 'Section 4.2', 'sec:evaluation': 'Section 4.3',
    'sec:exp_matrix': 'Section 4.4', 'sec:kan_head': 'Section 3.3',
}


# ============================================================
#  LATEX → UNICODE CONVERSION
# ============================================================

def _superscript(s):
    sup_map = {'0': '⁰', '1': '¹', '2': '²', '3': '³', '4': '⁴',
               '5': '⁵', '6': '⁶', '7': '⁷', '8': '⁸', '9': '⁹',
               '+': '⁺', '-': '⁻', '=': '⁼', '(': '⁽', ')': '⁾',
               'n': 'ⁿ', 'i': 'ⁱ'}
    return ''.join(sup_map.get(c, f'^{c}') for c in s)


def _clean_math(s):
    """Convert math-mode content to Unicode."""
    s = s.replace(r'\Rsq', 'R²')
    s = s.replace(r'R^{2}', 'R²').replace(r'R^2', 'R²')
    s = s.replace(r'\lamhole', 'λ_hole').replace(r'\lamcat', 'λ_cation')
    s = s.replace(r'\lamtrip', 'λ_triplet')
    s = s.replace(r'\lam', 'λ').replace(r'\lambda', 'λ')
    s = s.replace(r'\Delta', 'Δ').replace(r'\delta', 'δ')
    s = s.replace(r'\sigma', 'σ').replace(r'\mu', 'μ')
    s = s.replace(r'\theta', 'θ').replace(r'\phi', 'φ')
    s = s.replace(r'\pi', 'π').replace(r'\alpha', 'α').replace(r'\beta', 'β')
    s = s.replace(r'\gamma', 'γ').replace(r'\epsilon', 'ε')
    s = s.replace(r'\geq', '≥').replace(r'\leq', '≤')
    s = s.replace(r'\neq', '≠').replace(r'\approx', '≈')
    s = s.replace(r'\in', '∈').replace(r'\notin', '∉')
    s = s.replace(r'\pm', '±').replace(r'\mp', '∓')
    s = s.replace(r'\times', '×').replace(r'\cdot', '·')
    s = s.replace(r'\sim', '~').replace(r'\propto', '∝')
    s = s.replace(r'\sum', 'Σ').replace(r'\prod', 'Π')
    s = s.replace(r'\infty', '∞').replace(r'\partial', '∂')
    s = s.replace(r'\exp', 'exp').replace(r'\log', 'log')
    s = s.replace(r'\min', 'min').replace(r'\max', 'max')
    s = re.sub(r'\\AA\{?\}?', 'Å', s)
    s = s.replace('^{\\circ}', '°').replace(r'\circ', '°')
    # \text, \mathrm, etc.
    s = re.sub(r'\\text\{([^}]+)\}', r'\1', s)
    s = re.sub(r'\\mathrm\{([^}]+)\}', r'\1', s)
    s = re.sub(r'\\mathbf\{([^}]+)\}', r'\1', s)
    s = re.sub(r'\\mathbb\{([^}]+)\}', r'\1', s)
    s = re.sub(r'\\mathcal\{([^}]+)\}', r'\1', s)
    # Subscripts/superscripts
    s = re.sub(r'_\{([^}]+)\}', r'_\1', s)
    s = re.sub(r'\^\{([^}]+)\}', lambda m: _superscript(m.group(1)), s)
    s = re.sub(r'_([a-zA-Z0-9])', r'_\1', s)
    # Fractions
    s = re.sub(r'\\frac\{([^}]+)\}\{([^}]+)\}', r'(\1/\2)', s)
    s = re.sub(r'\\sqrt\{([^}]+)\}', r'√\1', s)
    # Delimiters
    s = s.replace(r'\|', '|').replace(r'\!', '')
    s = s.replace(r'\left', '').replace(r'\right', '')
    # Remaining commands
    s = re.sub(r'\\[a-zA-Z]+\{([^}]*)\}', r'\1', s)
    s = re.sub(r'\\[a-zA-Z]+', '', s)
    s = s.replace('\\', '')
    s = s.replace('{', '').replace('}', '')
    return s


def _convert_lists(text):
    """Convert enumerate/itemize to plain text."""
    def enum_replace(m):
        items = re.split(r'\\item\s*', m.group(1))
        items = [it.strip() for it in items if it.strip()]
        return '\n'.join(f'{i+1}. {it}' for i, it in enumerate(items))
    text = re.sub(r'\\begin\{enumerate\}(.*?)\\end\{enumerate\}',
                  enum_replace, text, flags=re.DOTALL)

    def item_replace(m):
        items = re.split(r'\\item\s*', m.group(1))
        items = [it.strip() for it in items if it.strip()]
        return '\n'.join(f'  \u2022 {it}' for it in items)
    text = re.sub(r'\\begin\{itemize\}(.*?)\\end\{itemize\}',
                  item_replace, text, flags=re.DOTALL)
    return text


def latex_to_unicode(text, cite_map):
    """Full LaTeX → Unicode conversion."""
    text = re.sub(r'(?<!\\)%.*$', '', text, flags=re.MULTILINE)

    # Custom commands (longer first)
    text = text.replace(r'\lamhole', 'λ_hole')
    text = text.replace(r'\lamcat', 'λ_cation')
    text = text.replace(r'\lamtrip', 'λ_triplet')
    text = text.replace(r'\thetaenc', 'θ_enc')
    text = text.replace(r'\thetahead', 'θ_head')
    text = text.replace(r'\Rsq', 'R²')
    text = text.replace(r'\lam', 'λ')
    for cmd in [r'\FedPer', r'\FedBN', r'\FedAvg']:
        name = cmd[1:]
        text = text.replace(cmd + '{}', name)
        text = text.replace(cmd, name)

    # Citations
    def replace_cite(m):
        keys = [k.strip() for k in m.group(1).split(',')]
        refs = [cite_map.get(k, k) for k in keys]
        return '[' + '; '.join(refs) + ']'
    text = re.sub(r'\\cite\{([^}]+)\}', replace_cite, text)

    # Cross-references with prefix
    def replace_fig_ref(m):
        prefix = m.group(1)
        label = m.group(2)
        display = REF_MAP.get(label, label)
        if prefix.rstrip('.~ ').lower() in display.lower():
            return display
        return f'{prefix} {display}'
    text = re.sub(
        r'(Fig\.\s*~?\s*|Table\s*~?\s*|Section\s*~?\s*|Eq\.\s*~?\s*)\\ref\{([^}]+)\}',
        replace_fig_ref, text)
    # Bare \ref
    text = re.sub(r'\\ref\{([^}]+)\}',
                  lambda m: REF_MAP.get(m.group(1), m.group(1)), text)

    # Equations
    def simplify_equation(m):
        eq = _clean_math(m.group(1))
        return f'\n[Equation: {eq}]\n'
    text = re.sub(
        r'\\begin\{equation\}(?:\s*\\label\{[^}]*\})?\s*(.*?)\s*\\end\{equation\}',
        simplify_equation, text, flags=re.DOTALL)

    # Inline math
    text = re.sub(r'\$([^$]+)\$', lambda m: _clean_math(m.group(1)), text)

    # Lists
    text = _convert_lists(text)

    # Formatting commands
    text = re.sub(r'\\textbf\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\\textit\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\\emph\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\\texttt\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\\textsc\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\\textsuperscript\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\\url\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\\label\{[^}]*\}', '', text)

    # Special chars
    text = text.replace(r'\AA{}', 'Å').replace(r'\AA', 'Å')
    text = text.replace('---', '\u2014')
    text = text.replace('--', '\u2013')
    text = text.replace(r'~', '\u00A0')
    text = text.replace(r'\,', ' ')
    text = text.replace(r'\ ', ' ')
    text = text.replace(r'\&', '&')
    text = text.replace(r'\%', '%')
    text = text.replace(r'\#', '#')
    text = text.replace(r'\_', '_')
    text = text.replace('{,}', ',')
    text = text.replace(r'\"u', 'ü').replace(r'\"o', 'ö').replace(r'\"a', 'ä')
    text = text.replace(r"\'e", 'é').replace(r"\^o", 'ô')

    # Remove braces
    text = text.replace('{', '').replace('}', '')

    # Remove remaining \commands
    text = re.sub(r'\\[a-zA-Z]+\b', '', text)

    # Clean whitespace
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()


# ============================================================
#  TEX PARSER
# ============================================================

def _extract_brace_content(s, start):
    """Extract content between matched braces starting at position `start`.
    Handles nested braces correctly.
    Example: _extract_brace_content(r'\\section{Foo \\FedPer{} + bar}\\label{x}', 8)
             → 'Foo \\FedPer{} + bar'
    """
    depth = 0
    i = start
    while i < len(s):
        if s[i] == '{':
            depth += 1
            if depth == 1:
                content_start = i + 1
        elif s[i] == '}':
            depth -= 1
            if depth == 0:
                return s[content_start:i]
        i += 1
    # Fallback: return everything after first {
    return s[start + 1:]


def parse_tex(tex_content, cite_map):
    """Parse LaTeX body into structured element list. Stops at appendix."""
    elements = []
    skip_preamble = True
    in_abstract = False
    in_table = False
    in_figure = False
    in_appendix = False
    buffer = []
    table_buf = []
    fig_buf = []

    sec_num = 0
    subsec_num = 0
    subsubsec_num = 0
    tab_counter = 0

    def flush_buffer():
        nonlocal buffer
        if buffer:
            raw = '\n'.join(buffer)
            text = latex_to_unicode(raw, cite_map)
            if text.strip():
                elements.append(('text', '', text.strip()))
            buffer = []

    for line in tex_content.split('\n'):
        stripped = line.strip()

        if skip_preamble:
            if r'\begin{document}' in stripped:
                skip_preamble = False
            continue

        if r'\end{document}' in stripped:
            break

        # Stop at appendix / clearpage before appendix figures
        if r'\appendix' in stripped:
            flush_buffer()
            in_appendix = True
            continue
        if in_appendix:
            continue

        if stripped in (r'\maketitle', r'\clearpage'):
            continue

        # Abstract
        if r'\begin{abstract}' in stripped:
            flush_buffer()
            in_abstract = True
            buffer = []
            continue
        if r'\end{abstract}' in stripped:
            in_abstract = False
            raw = '\n'.join(buffer)
            text = latex_to_unicode(raw, cite_map)
            elements.append(('abstract', 'Abstract', text.strip()))
            buffer = []
            continue
        if in_abstract:
            buffer.append(stripped)
            continue

        # Keywords
        if stripped.startswith(r'\textbf{Keywords'):
            flush_buffer()
            text = latex_to_unicode(stripped, cite_map)
            elements.append(('keywords', '', text))
            continue

        # Table environment → parse inline
        if r'\begin{table}' in stripped:
            flush_buffer()
            in_table = True
            table_buf = []
            tab_counter += 1
            continue
        if r'\end{table}' in stripped:
            in_table = False
            content = '\n'.join(table_buf)
            cap_match = re.search(r'\\caption\{(.+?)\}', content, re.DOTALL)
            cap = latex_to_unicode(cap_match.group(1), cite_map) if cap_match else 'Table'
            label_match = re.search(r'\\label\{(.+?)\}', content)
            label = label_match.group(1) if label_match else ''
            # Determine table type
            if 'subgroup' in label.lower() or 'subgroup' in cap.lower():
                elements.append(('table_subgroup', f'Table {tab_counter}', cap))
            else:
                elements.append(('table_placeholder', f'Table {tab_counter}', cap))
            table_buf = []
            continue
        if in_table:
            table_buf.append(stripped)
            continue

        # Figure environment → skip (we insert figures inline via FIGURE_PLACEMENTS)
        if r'\begin{figure}' in stripped:
            flush_buffer()
            in_figure = True
            fig_buf = []
            continue
        if r'\end{figure}' in stripped:
            in_figure = False
            fig_buf = []
            continue
        if in_figure:
            fig_buf.append(stripped)
            continue

        # Section headers
        sec_match = re.match(r'\\section\*?\{(.+)\}', stripped)
        if sec_match:
            flush_buffer()
            title_raw = _extract_brace_content(stripped, stripped.index('{'))
            is_starred = r'\section*' in stripped
            if not is_starred:
                sec_num += 1
                subsec_num = 0
                subsubsec_num = 0
                title = f'{sec_num}. {latex_to_unicode(title_raw, cite_map)}'
            else:
                title = latex_to_unicode(title_raw, cite_map)
            elements.append(('heading1', title, ''))
            continue

        subsec_match = re.match(r'\\subsection\*?\{(.+)\}', stripped)
        if subsec_match:
            flush_buffer()
            title_raw = _extract_brace_content(stripped, stripped.index('{'))
            is_starred = r'\subsection*' in stripped
            if not is_starred:
                subsec_num += 1
                subsubsec_num = 0
                title = f'{sec_num}.{subsec_num} {latex_to_unicode(title_raw, cite_map)}'
            else:
                title = latex_to_unicode(title_raw, cite_map)
            elements.append(('heading2', title, ''))
            continue

        subsubsec_match = re.match(r'\\subsubsection\*?\{(.+)\}', stripped)
        if subsubsec_match:
            flush_buffer()
            title_raw = _extract_brace_content(stripped, stripped.index('{'))
            subsubsec_num += 1
            title = f'{sec_num}.{subsec_num}.{subsubsec_num} {latex_to_unicode(title_raw, cite_map)}'
            elements.append(('heading3', title, ''))
            continue

        # \paragraph{Title}
        para_match = re.match(r'\\paragraph\{(.+?)\}(.*)', stripped)
        if para_match:
            flush_buffer()
            ptitle = latex_to_unicode(para_match.group(1), cite_map)
            remainder = para_match.group(2).strip()
            elements.append(('paragraph_head', ptitle, ''))
            if remainder:
                buffer.append(remainder)
            continue

        # Bibliography commands → skip
        if stripped.startswith(r'\bibliographystyle') or stripped.startswith(r'\bibliography'):
            flush_buffer()
            continue

        # Skip comments and labels
        if stripped.startswith('%') or stripped.startswith(r'\label'):
            continue

        # Regular text
        if stripped:
            buffer.append(stripped)
        elif buffer:
            flush_buffer()

    flush_buffer()
    return elements


# Tables to insert BEFORE specific headings
TABLE_PLACEMENTS = [
    ('5. ',  ['experiment_table']),   # Table 1 before Results section
    ('5.5 ', ['centralized_table']),  # Table 3 before Ablation section
]


def insert_figures_into_elements(elements):
    """Post-process: insert figure and table elements before specific headings."""
    result = []
    for elem in elements:
        etype, title, content = elem
        if etype in ('heading1', 'heading2', 'heading3'):
            # Insert tables before specific headings
            for prefix, tbl_keys in TABLE_PLACEMENTS:
                if title.startswith(prefix):
                    for tk in tbl_keys:
                        result.append(('insert_table', tk, ''))
                    break
            # Insert figures before specific headings
            for prefix, fig_keys in FIGURE_PLACEMENTS:
                if title.startswith(prefix):
                    for fk in fig_keys:
                        result.append(('figure', fk, ''))
                    break
        result.append(elem)
    return result


# ============================================================
#  TABLE CONSTRUCTION
# ============================================================

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def add_formatted_table(doc, title, headers, rows, col_widths=None):
    """Add a professionally formatted table."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(11)
    run.font.name = 'Times New Roman'

    n_cols = len(headers)
    table = doc.add_table(rows=len(rows) + 1, cols=n_cols)
    table.style = 'Table Grid'
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        run.font.name = 'Times New Roman'
        set_cell_shading(cell, 'D9E1F2')

    for i, row_data in enumerate(rows):
        for j, val in enumerate(row_data):
            cell = table.rows[i + 1].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(val))
            run.font.size = Pt(9)
            run.font.name = 'Times New Roman'

    if col_widths:
        for row in table.rows:
            for j, w in enumerate(col_widths):
                if j < len(row.cells):
                    row.cells[j].width = Inches(w)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    return table


def build_experiment_table(doc):
    """Table 1: Key experiments summary."""
    csv_path = RESULTS_DIR / 'all_experiments_summary.csv'
    if not csv_path.exists():
        p = doc.add_paragraph('[Table 1: all_experiments_summary.csv not found]')
        p.runs[0].font.color.rgb = RGBColor(200, 0, 0)
        return

    with open(csv_path, 'r') as f:
        reader = csv.DictReader(f)
        all_rows = list(reader)

    key_ids = ['E1', 'E2', 'E3', 'E7', 'E9', 'E10', 'E14', 'E15', 'E16',
               'E18', 'E30', 'E32', 'E40', 'E50', 'E51', 'E52']
    rows_data = []
    for r in all_rows:
        eid = r.get('exp_id', '')
        if eid in key_ids:
            def fmt(v):
                if not v or v.strip() == '':
                    return '\u2014'
                if '±' in v:
                    mean_part = v.split('±')[0]
                    try:
                        return f'{float(mean_part):.3f}'
                    except ValueError:
                        return v
                try:
                    return f'{float(v):.3f}'
                except ValueError:
                    return v

            rows_data.append([
                eid, r.get('encoder', ''), r.get('fed_strategy', 'None'),
                r.get('desc', ''),
                fmt(r.get('R2_A', '')), fmt(r.get('R2_B', '')),
                fmt(r.get('R2_C', '')), fmt(r.get('R2_Dh', '')),
                fmt(r.get('R2_Dt', '')),
            ])

    headers = ['ID', 'Encoder', 'Strategy', 'Description',
               'R²_A', 'R²_B', 'R²_C', 'R²_D(hole)', 'R²_D(trip)']
    add_formatted_table(doc,
                        'Table 1. Summary of Key Experiments (R² values)',
                        headers, rows_data,
                        col_widths=[0.4, 0.6, 0.7, 1.8, 0.6, 0.6, 0.6, 0.7, 0.7])


def build_subgroup_table(doc):
    """Table 2: Subgroup analysis."""
    headers = ['Subgroup', 'n', 'GIN R²', 'GIN MAE', 'SchNet R²', 'SchNet MAE']
    rows = [
        ['All molecules', '53', '\u22120.031', '0.366', '+0.181', '0.304'],
        ['Rigid High-DA (DA\u22650.25, rot=0)', '13', '+0.317', '0.258', '+0.852', '0.117'],
        ['Flexible (rot\u22651)', '12', '\u22120.363', '0.554', '+0.157', '0.388'],
    ]
    add_formatted_table(doc,
                        'Table 2. Subgroup Analysis of Client D LOOCV Results',
                        headers, rows,
                        col_widths=[2.2, 0.4, 0.7, 0.7, 0.7, 0.7])


def build_centralized_table(doc):
    """Table 3: Centralized vs Federated."""
    csv_path = RESULTS_DIR / 'centralized_upper_bound.csv'
    cent_hole_r2, cent_hole_mae = '\u22120.004', '0.300'
    cent_trip_r2, cent_trip_mae = '+0.148', '0.563'
    if csv_path.exists():
        with open(csv_path, 'r') as f:
            reader = csv.DictReader(f)
            for r in reader:
                if r.get('target') == 'hole':
                    cent_hole_r2 = f"{float(r['R2']):+.3f}"
                    cent_hole_mae = f"{float(r['MAE']):.3f}"
                elif r.get('target') == 'triplet':
                    cent_trip_r2 = f"{float(r['R2']):+.3f}"
                    cent_trip_mae = f"{float(r['MAE']):.3f}"

    headers = ['Method', 'R²(hole)', 'MAE(hole)', 'R²(triplet)', 'MAE(triplet)']
    rows = [
        ['Local SchNet+KAN (E50)', '\u22120.543', '0.389', '\u22120.258', '0.666'],
        ['Federated SchNet+KAN (E51)', '+0.181', '0.304', '\u22121.026', '0.848'],
        ['Centralized SchNet+KAN', cent_hole_r2, cent_hole_mae, cent_trip_r2, cent_trip_mae],
    ]
    add_formatted_table(doc,
                        'Table 3. Centralized vs. Federated Training on Client D',
                        headers, rows,
                        col_widths=[2.0, 0.9, 0.9, 0.9, 0.9])


# ============================================================
#  FIGURE INSERTION
# ============================================================

def add_figure(doc, fig_key):
    """Insert a figure with centered image + italic caption."""
    fdef = FIGURE_DEFS.get(fig_key)
    if not fdef:
        return
    img_path = FIGURES_DIR / fdef['file']
    width = fdef['width']
    caption = fdef['caption']

    if img_path.exists():
        # Image paragraph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(str(img_path), width=Inches(width))
    else:
        # Placeholder if file missing
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(f'[Image not found: {fdef["file"]}]')
        run.font.color.rgb = RGBColor(200, 0, 0)
        run.font.size = Pt(10)
        run.font.name = 'Times New Roman'

    # Caption paragraph
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(caption)
    run.italic = True
    run.font.size = Pt(10)
    run.font.name = 'Times New Roman'


# ============================================================
#  REFERENCE LIST
# ============================================================

def add_reference_list(doc, bib_entries):
    doc.add_paragraph('References', style='Heading 1')
    cite_order = [
        'Marcus1993', 'Coropceanu2007', 'Uoyama2012_TADF', 'Nelsen1987',
        'AtahanEvrenk2019_RE', 'Li2023_SchNet', 'McMahan2017_FedAvg',
        'FedChem2022', 'Heyndrickx2023', 'FLAP2023',
        'Xu2019_GIN', 'Schutt2018_SchNet', 'KAN2024',
        'Arivazhagan2019_FedPer', 'Li2021_FedBN',
        'Gilmer2017_MPNN', 'Gasteiger2020_DimeNet', 'Schutt2021_PaiNN',
        'RDKit2023', 'EfficientKAN2024',
        'You2020_GraphCL', 'Hu2020_SSL',
        'KAGNN2025', 'FedLG2025',
    ]
    remaining = [k for k in bib_entries if k not in cite_order]
    all_keys = cite_order + remaining

    for i, key in enumerate(all_keys, 1):
        if key not in bib_entries:
            continue
        entry = bib_entries[key]
        author = entry.get('author', 'Unknown')
        if ' and ' in author:
            authors = [a.strip() for a in author.split(' and ')]
            first = authors[0].split(',')[0].strip()
            if len(authors) > 2:
                author_str = f'{first} et al.'
            else:
                second = authors[1].split(',')[0].strip()
                author_str = f'{first} and {second}'
        elif author:
            author_str = author.split(',')[0].strip()
        else:
            author_str = key

        title = entry.get('title', '')
        journal = entry.get('journal', entry.get('booktitle', ''))
        year = entry.get('year', '')
        volume = entry.get('volume', '')
        pages = entry.get('pages', '')

        ref_text = f'[{i}] {author_str}. {title}. '
        if journal:
            ref_text += journal
            if volume:
                ref_text += f', {volume}'
            if pages:
                ref_text += f', {pages}'
            ref_text += f', {year}.'
        else:
            ref_text += f'{year}.'

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-1.0)
        run = p.add_run(ref_text)
        run.font.size = Pt(10)
        run.font.name = 'Times New Roman'


# ============================================================
#  DOCUMENT BUILDER
# ============================================================

def setup_styles(doc):
    """Configure Word styles."""
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rPr = normal._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:eastAsia'), 'Times New Roman')

    sizes = {1: 14, 2: 12, 3: 11}
    space_before = {1: 18, 2: 14, 3: 10}
    space_after = {1: 8, 2: 6, 3: 4}
    for level in [1, 2, 3]:
        style = doc.styles[f'Heading {level}']
        style.font.name = 'Times New Roman'
        style.font.bold = True
        style.font.size = Pt(sizes[level])
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(space_before[level])
        style.paragraph_format.space_after = Pt(space_after[level])
        style.paragraph_format.line_spacing = 1.5
        pPr = style._element.get_or_add_pPr()
        for numPr in pPr.findall(qn('w:numPr')):
            pPr.remove(numPr)


def build_docx(elements, bib_entries, cite_map):
    """Build the Word document from enriched element list."""
    doc = Document()
    setup_styles(doc)

    # ---- Title ----
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(
        'Personalized Federated Learning with Graph Neural Networks '
        'and Kolmogorov\u2013Arnold Networks for Cross-Domain '
        'Molecular Reorganization Energy Prediction'
    )
    run.bold = True
    run.font.size = Pt(16)
    run.font.name = 'Times New Roman'

    # ---- Authors ----
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run('Can Leng\u00b9, Author Two\u00b9')
    run.font.size = Pt(12)
    run.font.name = 'Times New Roman'

    for etype, title, content in elements:
        if etype == 'abstract':
            doc.add_paragraph(title, style='Heading 1')
            for para in content.split('\n\n'):
                para = para.strip()
                if para:
                    p = doc.add_paragraph(para)
                    p.style = doc.styles['Normal']
                    p.paragraph_format.first_line_indent = Inches(0.5)

        elif etype == 'keywords':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run(content)
            run.italic = True
            run.font.size = Pt(11)
            run.font.name = 'Times New Roman'

        elif etype == 'heading1':
            doc.add_paragraph(title, style='Heading 1')

        elif etype == 'heading2':
            doc.add_paragraph(title, style='Heading 2')

        elif etype == 'heading3':
            doc.add_paragraph(title, style='Heading 3')

        elif etype == 'paragraph_head':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            run = p.add_run(f'{title}. ')
            run.bold = True
            run.font.size = Pt(12)
            run.font.name = 'Times New Roman'

        elif etype == 'text':
            paras = content.split('\n\n')
            for para_text in paras:
                para_text = para_text.strip()
                if not para_text:
                    continue
                lines = para_text.split('\n')
                for line in lines:
                    line = line.strip()
                    if not line:
                        continue
                    p = doc.add_paragraph(line)
                    p.style = doc.styles['Normal']
                    if not re.match(r'^(\d+\.|•)', line):
                        p.paragraph_format.first_line_indent = Inches(0.5)

        elif etype == 'figure':
            add_figure(doc, title)  # title holds the fig_key

        elif etype == 'table_subgroup':
            build_subgroup_table(doc)

        elif etype == 'table_placeholder':
            pass  # Handled via TABLE_PLACEMENTS

        elif etype == 'insert_table':
            if title == 'experiment_table':
                build_experiment_table(doc)
            elif title == 'centralized_table':
                build_centralized_table(doc)

    # ---- References ----
    doc.add_page_break()
    add_reference_list(doc, bib_entries)

    return doc


# ============================================================
#  VERIFICATION
# ============================================================

def verify_docx(doc, docx_path):
    print('\n' + '=' * 60)
    print('VERIFICATION REPORT')
    print('=' * 60)

    total_paragraphs = len(doc.paragraphs)
    print(f'Total paragraphs: {total_paragraphs}')

    h1, h2, h3 = [], [], []
    for p in doc.paragraphs:
        if p.style.name == 'Heading 1':
            h1.append(p.text)
        elif p.style.name == 'Heading 2':
            h2.append(p.text)
        elif p.style.name == 'Heading 3':
            h3.append(p.text)

    print(f'\nHeading 1 ({len(h1)}):')
    for h in h1:
        print(f'  - {h}')
    print(f'\nHeading 2 ({len(h2)}):')
    for h in h2:
        print(f'  - {h}')
    print(f'\nHeading 3 ({len(h3)}):')
    for h in h3:
        print(f'  - {h}')

    print(f'\nTables: {len(doc.tables)}')
    for i, t in enumerate(doc.tables, 1):
        print(f'  Table {i}: {len(t.rows)} rows x {len(t.columns)} cols')

    # Count inline images
    img_count = 0
    for rel in doc.part.rels.values():
        if 'image' in rel.reltype:
            img_count += 1
    print(f'\nInline images: {img_count}')

    # Check for residual LaTeX
    residual_patterns = [r'\\[a-zA-Z]+', r'\$[^$]+\$', r'\\begin\{', r'\\end\{']
    residuals = []
    for p in doc.paragraphs:
        text = p.text
        for pat in residual_patterns:
            for m in re.finditer(pat, text):
                if m.group() not in ('\\n', '\\t', '\\r'):
                    residuals.append(m.group())

    if residuals:
        unique = set(residuals)
        print(f'\nResidual LaTeX ({len(residuals)} occurrences, {len(unique)} unique):')
        for r in sorted(unique):
            print(f'  WARNING: {r}')
    else:
        print('\nNo residual LaTeX commands found.')

    fsize = os.path.getsize(docx_path)
    print(f'\nFile size: {fsize / 1024:.1f} KB')
    print(f'Saved to: {docx_path}')
    print('=' * 60)


# ============================================================
#  MAIN
# ============================================================

if __name__ == '__main__':
    print('Reading references.bib...')
    bib_entries = parse_bib(BIB_PATH)
    cite_map = make_cite_map(bib_entries)
    print(f'  Parsed {len(bib_entries)} bib entries')

    print('Reading main.tex...')
    tex_content = TEX_PATH.read_text(encoding='utf-8')

    print('Parsing LaTeX structure...')
    elements = parse_tex(tex_content, cite_map)
    print(f'  Found {len(elements)} document elements')

    print('Inserting figures at inline positions...')
    elements = insert_figures_into_elements(elements)
    print(f'  Enriched to {len(elements)} elements (with figures)')

    print('Building Word document...')
    doc = build_docx(elements, bib_entries, cite_map)

    print(f'Saving to {DOCX_PATH}...')
    doc.save(str(DOCX_PATH))

    verify_docx(doc, DOCX_PATH)
    print('\nDone!')
